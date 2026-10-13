# Generates scripts/fixtures/real.xlsx — a "real" Excel file (shared strings, date-formatted cells,
# datetimes, numbers, a title row above the header) used by scripts/import.test.ts.
import datetime as dt
from pathlib import Path
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "日程"
ws.append(["2026 秋季日程导出"])
ws.append([])
ws.append(["任务名称", "开始日期", "截止日期", "提醒时间", "工时", "进度", "优先级"])
ws.append(["论文初稿", dt.date(2026, 10, 15), dt.date(2026, 10, 20), dt.datetime(2026, 10, 19, 9, 30), 3, 0.25, "高"])
ws.append(["小组讨论", dt.date(2026, 11, 2), None, None, "45min", None, "中"])
ws.append(["整理笔记", None, "2026/11/05", None, None, "10%", "低"])
for row in ws.iter_rows(min_row=4, min_col=2, max_col=3):
    for cell in row:
        cell.number_format = "yyyy-mm-dd"
ws["F4"].number_format = "0%"
wb.create_sheet("Empty")
Path(__file__).with_name("real.xlsx").parent.mkdir(parents=True, exist_ok=True)
wb.save(Path(__file__).with_name("real.xlsx"))
print("ok")
